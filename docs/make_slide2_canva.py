"""
docs/make_slide2_canva.py - rebuild slide 2 inside the team's Canva-exported deck
(20 x 11.25 in). Keeps the file's own header, footer and logos; replaces the body.

  hook : before -> after -> verified (real satellite photo, our 3D terrain,
         our heights next to airborne-LiDAR truth)
  then : 4 number cards, 3 problem -> answer cards, the pipeline arrows

    python docs/make_slide2_canva.py --deck <team slide-2 pptx> --lidar-img <lidar_check.jpg> --out <pptx>
"""
import argparse
import os
import re

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
NAVY, BLUE, RED = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x00, 0x70, 0xC0), RGBColor(0xC0, 0x39, 0x2B)
ORANGE, TEAL, GREEN = RGBColor(0xD9, 0x82, 0x2B), RGBColor(0x13, 0x8D, 0x75), RGBColor(0x2E, 0x7D, 0x32)
INK, GREY, LIGHT = RGBColor(0x1C, 0x1C, 0x1C), RGBColor(0x55, 0x5E, 0x6B), RGBColor(0xEE, 0xF3, 0xFA)
LINE, WHITE = RGBColor(0xC9, 0xD6, 0xE8), RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Arial"
KEEP = {"Group 2", "Group 4", "Group 7", "Group 10", "Group 13", "Group 15"}   # header, footer, logos


def rich(par, text, size, color, bold=False):
    for i, part in enumerate(re.split(r"\*\*", text)):
        if part:
            r = par.add_run()
            r.text = part
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), color
            r.font.bold = bold or (i % 2 == 1)


def text(s, x, y, w, h, paras, size, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=2):
    tf = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)).text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment, par.space_after = align, Pt(space)
        if isinstance(p, tuple):                      # (text, size, color, bold)
            rich(par, p[0], p[1], p[2], p[3])
        else:
            rich(par, p, size, color, bold)


def box(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=1.25):
    b = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        b.fill.background()
    else:
        b.fill.solid()
        b.fill.fore_color.rgb = fill
    if line is None:
        b.line.fill.background()
    else:
        b.line.color.rgb, b.line.width = line, Pt(lw)
    if radius is not None:
        b.adjustments[0] = radius
    b.shadow.inherit = False
    return b


def label(b, paras, size, color=WHITE, align=PP_ALIGN.CENTER):
    tf = b.text_frame
    tf.word_wrap, tf.vertical_anchor = True, MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.08)
    for i, p in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        rich(par, p, size if i == 0 else size - 3, color, bold=i == 0)


def picture(s, path, x, y, w, h):
    iw, ih = Image.open(path).size
    p = s.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    ar, ab = iw / ih, w / h
    if ar > ab:
        p.crop_left = p.crop_right = (1 - ab / ar) / 2
    else:
        p.crop_top = p.crop_bottom = (1 - ar / ab) / 2
    p.line.color.rgb, p.line.width = LINE, Pt(1)
    return p


def build(s, imgs):
    # ---- hook: before -> after -> verified
    iw, ih, aw, y = 5.5, 3.0, 0.75, 2.45
    steps = [(imgs["photo"], "1", "One satellite image", BLUE),
             (imgs["terrain"], "2", "Metric 3D terrain, in seconds", TEAL),
             (imgs["check"], "3", "Verified against airborne LiDAR", ORANGE)]
    for i, (img, n, cap, col) in enumerate(steps):
        x = 1.0 + i * (iw + aw)
        c = box(s, x, 1.93, 0.42, 0.42, fill=col, shape=MSO_SHAPE.OVAL)
        label(c, [n], 16)
        text(s, x + 0.5, 1.92, iw - 0.5, 0.45, cap, 18, color=col, bold=True, anchor=MSO_ANCHOR.MIDDLE)
        picture(s, img, x, y, iw, ih)
        if i < 2:
            a = box(s, x + iw + 0.08, y + ih / 2 - 0.42, aw - 0.16, 0.84, fill=NAVY, shape=MSO_SHAPE.RIGHT_ARROW)
            text(s, x + iw - 0.2, y + ih / 2 + 0.45, aw + 0.4, 0.4, "≈ 7 s" if i == 0 else "4.69 m",
                 12, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
    text(s, 1.0, y + ih + 0.02, 11.75, 0.32,
         "Chungthang, Sikkim: WorldView-2 image [9] → our DSM, rendered in our Three.js viewer",
         11, color=GREY, align=PP_ALIGN.LEFT)
    text(s, 13.5, y + ih + 0.02, 5.5, 0.32, "GAMUS test tile [3]: LiDAR truth vs our heights",
         11, color=GREY, align=PP_ALIGN.RIGHT)

    # ---- the numbers
    cards = [("4.69 m", "average height error (RMSE) on 40 held-out LiDAR test tiles [3]", BLUE),
             ("30% lower", "error than RS3DAda, a NeurIPS 2024 height model, same tiles [6]", TEAL),
             ("3.8 m", "full-DSM error on hilly terrain vs USGS airborne LiDAR [8]", ORANGE),
             ("≈ 7 s", "per 600 × 600 m scene on a laptop GPU, fully offline", GREEN)]
    for i, (v, c, col) in enumerate(cards):
        x = 1.0 + i * 4.567
        box(s, x, 5.92, 4.3, 1.3, fill=WHITE, line=col, lw=1.75)
        box(s, x, 5.92, 0.12, 1.3, fill=col)
        text(s, x + 0.25, 5.94, 4.0, 0.62, v, 30, color=col, bold=True)
        text(s, x + 0.25, 6.58, 4.0, 0.62, c, 13, color=GREY)

    # ---- problem -> our answer
    pa = [("LiDAR or stereo surveys: $5k–$100k+ and days of flying [2]",
           "One image you already have, turned into 3D in seconds"),
          ("Floods cost India USD 7.4 bn a year; maps are needed fast [1]",
           "A DSM GeoTIFF + 3D scene for flood and landslide screening"),
          ("AI depth models give only relative depth, not metres",
           "Heights in real metres, checked against LiDAR in the app")]
    for i, (p, a) in enumerate(pa):
        x = 1.0 + i * 6.1
        box(s, x, 7.42, 5.85, 1.3, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
        text(s, x + 0.2, 7.47, 5.5, 1.22,
             [(p, 12.5, RED, False), ("➜  **" + a + "**", 14, NAVY, False)],
             13, anchor=MSO_ANCHOR.MIDDLE, space=5)

    # ---- how it works (the team's arrow sequence)
    steps = [("Upload any image", "GeoTIFF, PNG, JPG or TIFF"),
             ("AI reads heights", "in metres, per pixel"),
             ("Add the terrain", "Copernicus / SRTM [7]"),
             ("Save the DSM", "GeoTIFF for GIS tools"),
             ("Fly through in 3D", "and check against LiDAR")]
    n, ov, w, x0 = 5, 0.18, 18.0, 1.0
    sw = (w + ov * (n - 1)) / n
    for i, (a, b) in enumerate(steps):
        c = box(s, x0 + i * (sw - ov), 8.92, sw, 1.25, fill=BLUE if i % 2 == 0 else NAVY,
                shape=MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON)
        label(c, [a, b], 16)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--deck", required=True)
    ap.add_argument("--lidar-img", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    tmp = os.path.dirname(os.path.abspath(a.out))
    photo = os.path.join(tmp, "s2_photo.jpg")
    Image.open(os.path.join(ROOT, "docs", "thumbs", "india_rgb.jpg")).save(photo, quality=94)
    terrain = os.path.join(tmp, "s2_terrain.jpg")
    Image.open(os.path.join(ROOT, "shots", "india_01_hero_dsm.jpg")).convert("RGB").crop(
        (640, 470, 1690, 1070)).save(terrain, quality=94)
    prs = Presentation(a.deck)
    s = prs.slides[0]
    for sh in list(s.shapes):
        if sh.name not in KEEP:
            sh._element.getparent().remove(sh._element)
    build(s, {"photo": photo, "terrain": terrain, "check": a.lidar_img})
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
