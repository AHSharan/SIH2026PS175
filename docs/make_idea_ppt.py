"""
docs/make_idea_ppt.py - the SIH 2026 idea deck (6 slides) built on the OFFICIAL
template. The template's design is kept as is (title style, SIH logo, blue
footer, slide numbers, title-slide artwork); only its placeholder text is
replaced and our content is added as native, editable shapes and charts.
Numbers: 40-tile GAMUS test of the 1,500-tile model (team's Colab run: RMSE 4.69 m,
MAE 2.29 m, r 0.81) and RS3DAda / predict-0 on the same tiles (RESULTS.md);
full-pipeline LiDAR results from results/lidar_benchmark.md.

    python docs/make_idea_ppt.py --official SIH2026-IDEA-Presentation-Format.pptx
        --team-deck <our earlier deck .pptx, for the ParallaX and SIT logos>
        --out ParallaX_DepthWizard_SIH2026.pptx
"""
import argparse
import os
import re
import tempfile
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

NAVY, BLUE, RED = RGBColor(0x1F, 0x38, 0x64), RGBColor(0x00, 0x70, 0xC0), RGBColor(0xC0, 0x39, 0x2B)
ORANGE, TEAL, GREEN = RGBColor(0xD9, 0x82, 0x2B), RGBColor(0x13, 0x8D, 0x75), RGBColor(0x2E, 0x7D, 0x32)
INK, GREY, LIGHT = RGBColor(0x1C, 0x1C, 0x1C), RGBColor(0x55, 0x5E, 0x6B), RGBColor(0xEE, 0xF3, 0xFA)
LINE, WHITE, SOFT = RGBColor(0xC9, 0xD6, 0xE8), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xB8, 0xC2, 0xCF)
FONT = "Arial"


# ----------------------------------------------------------------- assets
def prepare_assets(team_deck, d):
    """Logos from our earlier deck + cropped real outputs from shots/ and docs/thumbs/."""
    os.makedirs(d, exist_ok=True)
    with zipfile.ZipFile(team_deck) as z:
        z.extract("ppt/media/image6.png", d)          # ParallaX logo
        z.extract("ppt/media/image7.png", d)          # SIT logo strip (emblem on the left)
    Image.open(os.path.join(d, "ppt/media/image6.png")).save(os.path.join(d, "parallax.png"))
    Image.open(os.path.join(d, "ppt/media/image7.png")).crop((0, 0, 112, 146)).save(
        os.path.join(d, "emblem.png"))
    s = os.path.join(ROOT, "shots")
    Image.open(os.path.join(s, "india_01_hero_dsm.jpg")).convert("RGB").crop(
        (580, 430, 1720, 1080)).save(os.path.join(d, "hero.jpg"), quality=92)
    for n in ["india_03_buildings_floors", "india_04_flood", "india_02_height"]:
        Image.open(os.path.join(s, n + ".jpg")).convert("RGB").crop(
            (290, 150, 1345, 640)).save(os.path.join(d, n + ".jpg"), quality=92)
    for n in ["india_rgb", "india_tiles", "india_features", "india_ndsm", "india_dsm"]:
        Image.open(os.path.join(ROOT, "docs", "thumbs", n + ".jpg")).save(os.path.join(d, n + ".jpg"))
    return d


# ----------------------------------------------------------------- helpers
def remove(shape):
    shape._element.getparent().remove(shape._element)


def rich(par, text, size, color, bold=False, italic=False):
    """Add text to a paragraph; **...** marks bold runs."""
    for i, part in enumerate(re.split(r"\*\*", text)):
        if not part:
            continue
        r = par.add_run()
        r.text = part
        f = r.font
        f.name, f.size, f.color.rgb = FONT, Pt(size), color
        f.bold = bold or (i % 2 == 1)
        f.italic = italic


def text(slide, x, y, w, h, paras, size=12, color=INK, bold=False, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, space=4, italic=False, bullet=None):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, p in enumerate(paras if isinstance(paras, list) else [paras]):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        par.space_after = Pt(space)
        if bullet:
            rich(par, bullet + " ", size, BLUE, bold=True)
        rich(par, p, size, color, bold=bold, italic=italic)
    return tb


def box(slide, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, radius=None, lw=1.0):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(lw)
    if radius is not None:
        s.adjustments[0] = radius
    s.shadow.inherit = False
    s.text_frame.margin_left = s.text_frame.margin_right = Inches(0.06)
    return s


def label(s, paras, size=11, color=WHITE, bold=True, align=PP_ALIGN.CENTER):
    tf = s.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    for i, p in enumerate(paras):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align
        sz, b = (size, bold) if i == 0 else (size - 2, False)
        rich(par, p, sz, color, bold=b)


def picture(slide, path, x, y, w, h):
    """Place an image filling the box, cropped (not stretched) to its aspect."""
    iw, ih = Image.open(path).size
    p = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
    ar, ab = iw / ih, w / h
    if ar > ab:
        c = (1 - ab / ar) / 2
        p.crop_left = p.crop_right = c
    else:
        c = (1 - ar / ab) / 2
        p.crop_top = p.crop_bottom = c
    p.line.color.rgb = LINE
    p.line.width = Pt(0.75)
    return p


def section(slide, x, y, w, title, color):
    box(slide, x, y + 0.03, 0.08, 0.3, fill=color)
    text(slide, x + 0.14, y - 0.02, w, 0.4, title, size=14.5, color=color, bold=True)


def kpi(slide, x, y, w, h, value, caption, color):
    box(slide, x, y, w, h, fill=WHITE, line=color, lw=1.25)
    box(slide, x, y, 0.09, h, fill=color)
    text(slide, x + 0.18, y + 0.04, w - 0.25, 0.5, value, size=24, color=color, bold=True)
    text(slide, x + 0.18, y + 0.52, w - 0.25, h - 0.55, caption, size=10, color=GREY)


def chevrons(slide, x, y, w, h, steps):
    n, ov = len(steps), 0.12
    sw = (w + ov * (n - 1)) / n
    for i, (a, b) in enumerate(steps):
        shp = MSO_SHAPE.PENTAGON if i == 0 else MSO_SHAPE.CHEVRON
        s = box(slide, x + i * (sw - ov), y, sw, h, fill=BLUE if i % 2 == 0 else NAVY, shape=shp)
        s.text_frame.margin_left = Inches(0.28 if i else 0.12)
        s.text_frame.margin_right = Inches(0.2)
        label(s, [a, b], size=12.5)


def chart(slide, x, y, w, h, title, cats, series, colors):
    data = CategoryChartData()
    data.categories = cats
    for name, vals in series:
        data.add_series(name, vals)
    gf = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(x), Inches(y),
                                Inches(w), Inches(h), data)
    c = gf.chart
    c.has_title = True
    c.chart_title.text_frame.text = title
    tp = c.chart_title.text_frame.paragraphs[0]
    tp.runs[0].font.size, tp.runs[0].font.bold = Pt(12), True
    tp.runs[0].font.color.rgb, tp.runs[0].font.name = NAVY, FONT
    c.has_legend = True
    c.legend.position = XL_LEGEND_POSITION.BOTTOM
    c.legend.include_in_layout = False
    c.legend.font.size, c.legend.font.name = Pt(9.5), FONT
    va = c.value_axis
    va.has_major_gridlines = False
    va.visible = False
    ca = c.category_axis
    ca.tick_labels.font.size, ca.tick_labels.font.name = Pt(10), FONT
    ca.format.line.color.rgb = SOFT
    pl = c.plots[0]
    pl.gap_width, pl.overlap = 55, -5
    pl.has_data_labels = True
    dl = pl.data_labels
    dl.number_format, dl.number_format_is_linked = "0.00", False
    dl.position = XL_LABEL_POSITION.OUTSIDE_END
    dl.font.size, dl.font.name = Pt(9), FONT
    for s, col in zip(pl.series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = col
    return c


def team_logo(slide, a):
    slide.shapes.add_picture(os.path.join(a, "parallax.png"), Inches(0.3), Inches(0.36), Inches(2.15))


# ----------------------------------------------------------------- slides
def title_slide(s, a):
    tb = next(sh for sh in s.shapes if sh.name == "TextBox 9")
    lines = ["Problem Statement ID – SIH26175",
             "Problem Statement Title – DepthWizard: Single-View Height Estimation & 3D Flythrough",
             "Theme – Disaster Management", "PS Category – Software", "Team ID – T171",
             "Team Name (Registered on portal) – ParallaX"]
    paras = [p for p in tb.text_frame.paragraphs if p.runs]
    for p, t in zip(paras, lines):
        p.runs[0].text = t
        p.runs[0].font.size = Pt(20)                 # the full PS title needs two lines
        for r in p.runs[1:]:
            r.text = ""
        p.alignment = PP_ALIGN.LEFT                  # justified text spread the words apart
    s.shapes.add_picture(os.path.join(a, "emblem.png"), Inches(0.22), Inches(0.12), height=Inches(1.25))
    s.shapes.add_picture(os.path.join(a, "parallax.png"), Inches(10.8), Inches(1.3), Inches(2.3))


def slide_solution(s, a):
    # keep the template's own title run (its serif font); only the words change.
    # The placeholder has an empty first line and "IDEA TITLE" on the second.
    tf = s.shapes.title.text_frame
    idea = next(p for p in tf.paragraphs if "IDEA" in "".join(r.text for r in p.runs))
    idea.runs[0].text = "ONE IMAGE → 3D TERRAIN"   # must fit between the two logos
    idea.runs[0].font.size = Pt(30)
    for r in idea.runs[1:]:
        r._r.getparent().remove(r._r)
    for br in idea._p.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}br"):
        idea._p.remove(br)                   # the template breaks the line before its title
    for p in list(tf.paragraphs):
        if p._p is not idea._p:              # compare the XML, not the wrapper objects
            p._p.getparent().remove(p._p)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE    # same height as the other slide titles
    section(s, 0.35, 1.32, 4.1, "THE PROBLEM", RED)
    text(s, 0.4, 1.75, 4.05, 2.2, [
        "India loses about **USD 7.4 billion a year to floods** alone. [1]",
        "Flood and landslide maps need an elevation model, but today that means **LiDAR, "
        "stereo or radar**: $5k–$100k+ per survey and days of flying. [2]",
        "Depth AI trained on street photos gives only **relative depth**, not real metres from above."],
        size=12.5, bullet="▪", space=7)
    section(s, 4.65, 1.32, 4.1, "OUR SOLUTION", BLUE)
    text(s, 4.7, 1.75, 4.1, 2.3, [
        "**One satellite or aerial image in, a measured elevation model out**: a DSM GeoTIFF "
        "and a 3D scene you can fly through.",
        "Built on DINOv3-SAT [4] and trained on LiDAR heights [3], our model reads "
        "**height above ground in metres**; a 30 m terrain map adds the ground [7].",
        "GeoTIFF gives an absolute DSM; PNG/JPG a relative DSM, or a full DSM with 3 map points.",
        "It **checks itself**: load LiDAR or any reference and get RMSE, MAE, correlation and an error map."],
        size=11, bullet="▪", space=4)
    ui = os.path.join(a, "ui.jpg")
    picture(s, ui if os.path.exists(ui) else os.path.join(a, "hero.jpg"), 8.98, 1.36, 4.0, 2.28)
    text(s, 8.98, 3.66, 4.0, 0.3, ("Our app" if os.path.exists(ui) else "Our output") +
         ": Chungthang, Sikkim, from one WorldView-2 image [9]",
         size=9, color=GREY, italic=True, align=PP_ALIGN.CENTER)
    for i, (v, c, col) in enumerate([
            ("4.69 m", "average height error (RMSE) on 40 held-out airborne-LiDAR test tiles [3]", BLUE),
            ("30% lower", "error than RS3DAda, a NeurIPS 2024 height model, on the same tiles [6]", TEAL),
            ("3.8 m", "full-DSM error on hilly terrain against USGS airborne LiDAR; terrain map "
             "alone: 5.9 m [8]", ORANGE),
            ("≈ 7 s", "per 600 × 600 m scene on a laptop GPU, fully offline", GREEN)]):
        kpi(s, 0.35 + i * 3.19, 4.08, 3.07, 1.02, v, c, col)
    text(s, 0.35, 5.2, 12.6, 0.35, ["**WHAT IS NEW**   ·   heights come out in real metres directly, "
         "not as relative depth, and every result can be scored against LiDAR inside the app"],
         size=11.5, color=NAVY)
    chevrons(s, 0.35, 5.58, 12.63, 1.12, [
        ("Upload any image", "GeoTIFF, PNG, JPG or TIFF"),
        ("AI reads heights", "in metres, per pixel"),
        ("Add the terrain", "Copernicus or SRTM 30 m [7]"),
        ("Save the DSM", "GeoTIFF for QGIS / ArcGIS"),
        ("Fly through in 3D", "and check against LiDAR")])


def slide_technical(s, a):
    # the three key milestones, in the problem statement's own words
    cards = [
        ("1  ELEVATION EXTRACTION", BLUE, [
            "Frozen **DINOv3-SAT ViT-L/16** encoder, pre-trained by Meta on 493 M satellite images [4]",
            "Features from blocks 6, 12, 18 and 24 fused by a **DPT decoder** that we trained [5]",
            "Output: **height above ground in metres** for every pixel (nDSM)",
            "Any pixel size normalised to 0.5 m; 608 px tiles, 25% overlap, Gaussian blending, "
            "flip averaging"]),
        ("2  SCALE CALIBRATION", TEAL, [
            "Metric scale **learned from airborne-LiDAR heights** (masked Huber loss in metres) [3]",
            "**DSM = terrain + heights**: Copernicus GLO-30 or SRTM resampled onto the image grid, "
            "EGM2008 [7]",
            "PNG/JPG: relative DSM, or **3+ GCPs** fit an affine to UTM for a full metric DSM",
            "With reference data: **robust offset + scale** (RANSAC + Huber IRLS), tested on held-out pixels"]),
        ("3  VISUALIZATION LAYER", ORANGE, [
            "**Three.js / WebGL** terrain mesh up to 1024 × 1024 vertices, photo draped pixel-exact [10]",
            "Orbit and **first-person flight**; height, slope and contour shading; height profiles",
            "**Validation**: upload a reference and get RMSE, MAE, r, an error map and a side-by-side swipe",
            "Exports **GeoTIFF and GLB**; runs offline in the browser, no cloud"])]
    for i, (title, col, pts) in enumerate(cards):
        x = 0.35 + i * 4.265
        box(s, x, 1.3, 4.1, 2.5, fill=WHITE, line=col, lw=1.25)
        h = box(s, x, 1.3, 4.1, 0.4, fill=col)
        label(h, [title], size=12.5, align=PP_ALIGN.LEFT)
        text(s, x + 0.08, 1.76, 3.95, 2.0, pts, size=10.5, color=INK, bullet="▪", space=4)
    # the same steps as real outputs of our pipeline
    steps = [("india_rgb.jpg", "Input image", "WorldView-2, 0.31 m"),
             ("india_tiles.jpg", "Tiling", "608 px, 25% overlap"),
             ("india_features.jpg", "What the AI sees", "DINOv3-SAT features"),
             ("india_ndsm.jpg", "Heights (nDSM)", "metres above ground"),
             ("india_dsm.jpg", "+ terrain = DSM", "Copernicus 30 m"),
             ("hero.jpg", "3D flythrough", "Three.js, offline")]
    iw, gap, y0 = 1.93, (12.63 - 6 * 1.93) / 5, 3.92
    for i, (f, t, sub) in enumerate(steps):
        x = 0.35 + i * (iw + gap)
        picture(s, os.path.join(a, f), x, y0, iw, 1.18)
        c = box(s, x + 0.05, y0 + 0.05, 0.3, 0.3, fill=NAVY, line=WHITE, shape=MSO_SHAPE.OVAL, lw=1)
        label(c, [str(i + 1)], size=10)
        text(s, x, y0 + 1.2, iw, 0.26, t, size=10.5, color=NAVY, bold=True, align=PP_ALIGN.CENTER)
        text(s, x, y0 + 1.44, iw, 0.26, sub, size=9, color=GREY, align=PP_ALIGN.CENTER)
        if i < 5:
            ar = box(s, x + iw + gap / 2 - 0.07, y0 + 0.47, 0.14, 0.24, fill=BLUE,
                     shape=MSO_SHAPE.ISOSCELES_TRIANGLE)
            ar.rotation = 90
    chips = ["Python", "PyTorch", "DINOv3-SAT", "HF Transformers", "Rasterio · GDAL",
             "Copernicus GLO-30", "Three.js · WebGL", "GAMUS", "USGS 3DEP LiDAR"]
    widths = [max(0.9, 0.28 + 0.08 * len(c)) for c in chips]
    cg = (12.63 - sum(widths)) / (len(chips) - 1)
    x = 0.35
    for c, w in zip(chips, widths):
        b = box(s, x, 5.7, w, 0.32, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        label(b, [c], size=10, color=NAVY)
        x += w + cg
    b = box(s, 0.35, 6.14, 12.63, 0.66, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
    label(b, ["**TRAINING**   GAMUS [3]: 1,500 training, 48 validation and 40 held-out test tiles of "
              "aerial RGB with airborne-LiDAR heights  ·  AdamW, masked Huber loss in metres  ·  "
              "model chosen on the validation split only"], size=10.5, color=NAVY, bold=False)


def slide_feasibility(s, a):
    for i, (t, d, col) in enumerate([
            ("Working prototype today", "Upload an image and get a DSM GeoTIFF and a 3D scene in "
             "seconds. Covered by 101 automated tests.", BLUE),
            ("Checked against real LiDAR", "40 held-out GAMUS test tiles [3] and 8 USGS "
             "airborne-LiDAR sites [8], from towns to hills.", TEAL),
            ("Low cost, runs offline", "A laptop GPU and free data: GAMUS [3] and the Copernicus "
             "DEM [7]. No survey flights needed.", GREEN)]):
        x = 0.35 + i * 4.24
        box(s, x, 1.3, 4.1, 0.98, fill=WHITE, line=col, lw=1.25)
        box(s, x, 1.3, 0.09, 0.98, fill=col)
        text(s, x + 0.2, 1.33, 3.8, 0.35, t, size=13, color=col, bold=True)
        text(s, x + 0.2, 1.66, 3.8, 0.6, d, size=10, color=GREY)
    chart(s, 0.35, 2.4, 6.2, 2.2, "Error on 40 held-out LiDAR test tiles, metres [3]",
          ["RMSE", "MAE"],
          [("Predict 0 m", (9.35, 5.10)), ("RS3DAda [6]", (6.74, 3.47)), ("Ours", (4.69, 2.29))],
          [SOFT, RGBColor(0x7C, 0x8E, 0xA8), RED])
    text(s, 0.35, 4.58, 6.2, 0.28, "Correlation with LiDAR: **ours r = 0.81**, RS3DAda r = 0.60",
         size=10, color=NAVY, align=PP_ALIGN.CENTER)
    chart(s, 6.78, 2.4, 6.2, 2.2, "Full DSM vs airborne LiDAR, RMSE in metres [8]",
          ["Residential", "Sparse", "Hilly", "Forest"],
          [("Terrain map alone", (6.95, 4.99, 5.88, 8.43)), ("Ours: terrain + model", (5.40, 3.57, 3.84, 7.95))],
          [RGBColor(0x9F, 0xB6, 0xCD), RED])
    text(s, 6.78, 4.58, 6.2, 0.28, "Whole pipeline, GeoTIFF in → DSM out, on 8 USGS sites",
         size=10, color=NAVY, align=PP_ALIGN.CENTER)
    rows = [("Risk", "Mitigation"),
            ("Very tall towers and dense forest canopy",
             "More tall-building and forest tiles in training; per-landscape scores tracked on every run."),
            ("New sensors and regions (ISRO imagery)",
             "Pixel size normalised (0.3–1 m tested); the built-in reference check scores and calibrates on ISRO data."),
            ("Absolute scale from a single image",
             "Metric training on LiDAR heights, DEM / SRTM fusion, GCPs and reference-based calibration."),
            ("Angled (oblique) photos",
             "Built for top-down satellite and aerial imagery, as the problem statement specifies.")]
    tb = s.shapes.add_table(len(rows), 2, Inches(0.35), Inches(4.95), Inches(12.63), Inches(1.85)).table
    tb.columns[0].width, tb.columns[1].width = Inches(4.2), Inches(8.43)
    for r, (c1, c2) in enumerate(rows):
        for c, t in enumerate((c1, c2)):
            cell = tb.cell(r, c)
            cell.text = ""
            cell.margin_left = cell.margin_right = Inches(0.08)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            rich(cell.text_frame.paragraphs[0], t, 11 if r == 0 else 10, WHITE if r == 0 else INK, bold=r == 0)
            cell.fill.solid()
            cell.fill.fore_color.rgb = BLUE if r == 0 else (WHITE if r % 2 else LIGHT)
        tb.rows[r].height = Inches(0.33 if r == 0 else 0.38)


def slide_impact(s, a):
    section(s, 0.35, 1.28, 6, "WHO IT HELPS", BLUE)
    for i, (f, t, d) in enumerate([
            ("india_04_flood.jpg", "Disaster management",
             "Flood and landslide screening from one image. In Chungthang [9], a 36 m river rise "
             "reaches 12 of 204 buildings."),
            ("india_03_buildings_floors.jpg", "Urban planning",
             "Building heights and floor estimates, slopes and contours for growth, drainage and rooftop planning."),
            ("hero.jpg", "Remote and border areas",
             "Terrain where no LiDAR exists, such as Himalayan towns, for relief and mission planning.")]):
        x = 0.35 + i * 4.24
        box(s, x, 1.68, 4.1, 2.62, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
        picture(s, os.path.join(a, f), x + 0.1, 1.78, 3.9, 1.5)
        text(s, x + 0.15, 3.32, 3.85, 0.32, t, size=13, color=NAVY, bold=True)
        text(s, x + 0.15, 3.63, 3.85, 0.66, d, size=10, color=INK)
    section(s, 0.35, 4.42, 6, "BENEFITS", BLUE)
    for i, (t, d) in enumerate([
            ("ECONOMIC", "No survey flight: uses imagery already collected and a free 30 m terrain map."),
            ("SOCIAL", "Faster hazard maps for places that have never been surveyed."),
            ("ENVIRONMENTAL", "No extra aircraft or drone flights."),
            ("OPERATIONAL", "Offline on a laptop in seconds; standard GeoTIFF for existing GIS tools.")]):
        b = box(s, 0.35 + i * 3.19, 4.8, 3.07, 1.28, fill=BLUE if i % 2 == 0 else NAVY,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.12)
        label(b, [t, d], size=13)
    text(s, 0.35, 6.18, 12.63, 0.6, "“ParallaX turns images India already collects into measured 3D "
         "terrain in seconds, and shows its own error against LiDAR.”", size=13, color=NAVY,
         bold=True, italic=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def slide_references(s, a):
    refs = [("India flood losses", "UNISDR Global Assessment Report 2015: India AAL ≈ USD 9.8 bn/yr, floods USD 7.4 bn."),
            ("Elevation survey cost", "Kapcher / Candrone: LiDAR survey $5k–$100k+; Drone Launch Academy: survey-grade LiDAR over $100k."),
            ("GAMUS dataset", "Xiong et al., “GAMUS: Geometry-aware Multi-modal Semantic Segmentation Benchmark”, arXiv:2305.14914."),
            ("DINOv3", "Meta AI, 2025: self-supervised vision model; SAT-493M satellite weights (our frozen encoder)."),
            ("DPT head", "Ranftl et al., “Vision Transformers for Dense Prediction”, ICCV 2021."),
            ("RS3DAda baseline", "Song et al., “SynRS3D”, NeurIPS 2024, arXiv:2406.18151: public weights, run by us."),
            ("Terrain map", "Copernicus DEM GLO-30 (ESA / Airbus), 30 m; SRTM 30 m also supported."),
            ("Independent LiDAR check", "USGS 3D Elevation Program airborne LiDAR + USDA NAIP imagery (public domain)."),
            ("Indian test image", "Maxar Open Data Program, WorldView-2, Chungthang, Sikkim (CC BY-NC 4.0)."),
            ("Rendering", "Three.js: WebGL library for the interactive 3D flythrough, bundled for offline use."),
            ("Problem statement", "ISRO SAC, SIH 2026 PS 26175; github.com/IMG-PROCESS-SAC/SIH-DepthWizard-2026"),
            ("Our code and results", "github.com/AHSharan/SIH2026PS175: RESULTS.md, results/lidar_benchmark.md")]
    for i, (t, d) in enumerate(refs):
        col, row = i % 2, i // 2
        x, y = 0.35 + col * 6.38, 1.3 + row * 0.925
        box(s, x, y, 6.25, 0.84, fill=LIGHT, line=LINE)
        n = box(s, x, y, 0.55, 0.84, fill=BLUE)
        label(n, [str(i + 1)], size=15)
        text(s, x + 0.65, y + 0.06, 5.5, 0.3, t, size=11.5, color=NAVY, bold=True)
        text(s, x + 0.65, y + 0.36, 5.5, 0.48, d, size=9.5, color=GREY)


# ----------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--official", required=True)
    ap.add_argument("--team-deck", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--ui-shot", default=None,
                    help="optional screenshot of the running app, shown on slide 2")
    a = ap.parse_args()
    assets = prepare_assets(a.team_deck, os.path.join(tempfile.gettempdir(), "dw_ppt_assets"))
    ui = os.path.join(assets, "ui.jpg")
    if a.ui_shot:
        Image.open(a.ui_shot).convert("RGB").save(ui, quality=95)
    elif os.path.exists(ui):
        os.remove(ui)
    prs = Presentation(a.official)
    # the instructions slide (7) says to delete it before submitting
    ids = prs.slides._sldIdLst
    last = ids[6]
    prs.part.drop_rel(last.rId)
    ids.remove(last)
    sl = list(prs.slides)
    # slide 1 (title) is left as the official template for the team to fill in
    for s in sl[1:6]:
        for sh in list(s.shapes):
            if sh.name.startswith("Oval") or sh.name == "TextBox 8":
                remove(sh)
        team_logo(s, assets)
    slide_solution(sl[1], assets)
    slide_technical(sl[2], assets)
    slide_feasibility(sl[3], assets)
    slide_impact(sl[4], assets)
    slide_references(sl[5], assets)
    prs.save(a.out)
    print("wrote", a.out)


if __name__ == "__main__":
    main()
